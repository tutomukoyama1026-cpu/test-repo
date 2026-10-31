"""工程表と図面・計画書リストの日付のリンク（ずれの確認）。
使い方: python3 -I koutei_link.py 工程表.xlsx 工程リンク.xlsx 図面・計画書リスト.xlsx [--all]
- 対応表（工程リンク.xlsx の「対応表」シート、5行目以降）の1行ごとに、工程表のバー（行・名称）を探し、
  開始日（基準が終了日なら終了日）＋ずらし日数を出す。休みの日（工程表4行目のオレンジ、範囲外は土日）なら前の稼働日にする。
- リストの日付：①② AG 施工予定日、③ AD 現場施工予定日（数式なら保存されている値）。
- 出力（タブ区切り）：
    ずれ      シート 整理番号 名称 リストの日付 工程表から出した日付 差（リスト−工程表、日） 工程表の行・名称
    未設定    リストに日付が無いのに工程表に結びついているもの
    バーなし  対応表の行・名称のバーが工程表に見つからないもの（工程表を直したら対応表も直す）
    対応表なし リストにあって対応表に無い項目（対応表に行を足す）
  --all を付けると一致したもの（一致）と工程表なし（工程表なし）も出す。
"""
import sys, re, datetime, importlib.util, pathlib
import openpyxl

spec = importlib.util.spec_from_file_location('koutei', pathlib.Path(__file__).with_name('koutei.py'))
koutei = importlib.util.module_from_spec(spec); spec.loader.exec_module(koutei)

kpath, tpath, lpath = sys.argv[1:4]
show_all = '--all' in sys.argv[4:]
norm = lambda s: re.sub(r'[\s　]+', '', str(s or ''))

kws = openpyxl.load_workbook(kpath).worksheets[0]
bars = koutei.read_bars(kws)
hol = koutei.holidays(kws)
first, last = koutei.dt(koutei.C0), koutei.dt(koutei.C1)
is_off = lambda d: d in hol if first <= d <= last else d.weekday() >= 5

def find(row, name):
    hit = [b for b in bars if b[4] == row and norm(b[3]) == norm(name)]
    if not hit:  # 名称が少し変わったとき（前後の追記など）
        hit = [b for b in bars if b[4] == row and (norm(name) in norm(b[3]) or norm(b[3]) in norm(name))]
    return sorted(hit)

SHEETS = {'①': ('①施工図リスト', 11, 'AG'), '②': ('②製作図リスト', 11, 'AG'), '③': ('③計画書要領書リスト', 7, 'AD')}
lwb = openpyxl.load_workbook(lpath, data_only=True)
items = {}
for sh, (title, r0, col) in SHEETS.items():
    ws = lwb[title]
    for r in range(r0, ws.max_row + 1):
        no = ws.cell(r, 2).value
        if not no: continue
        v = ws[f'{col}{r}'].value
        items[(sh, str(no).strip())] = (ws.cell(r, 3).value, v.date() if isinstance(v, datetime.datetime) else v)

tws = openpyxl.load_workbook(tpath)['対応表']
done = set(); out = []
for r in range(5, tws.max_row + 1):
    sh, no, _, row, name, base, off, _memo = [tws.cell(r, c).value for c in range(1, 9)]
    if not (sh and no): continue
    key = (sh, str(no).strip()); done.add(key)
    lname, ldate = items.get(key, ('（リストに無し）', None))
    if norm(name) == '工程表なし':
        if show_all: out.append(('工程表なし', sh, no, lname, ldate or '', '', '', ''))
        continue
    hit = find(int(row), name)
    if not hit:
        out.append(('バーなし', sh, no, lname, ldate or '', '', '', f'行{row} {name}')); continue
    d = (hit[-1][1] if base == '終了日' else hit[0][0]) + datetime.timedelta(int(off or 0))
    while is_off(d): d -= datetime.timedelta(1)
    where = f'行{row} {hit[0][3]}'
    if not isinstance(ldate, datetime.date):
        out.append(('未設定', sh, no, lname, '', d, '', where))
    elif ldate != d:
        out.append(('ずれ', sh, no, lname, ldate, d, (ldate - d).days, where))
    elif show_all:
        out.append(('一致', sh, no, lname, ldate, d, 0, where))
for key, (lname, ldate) in items.items():
    if key not in done: out.append(('対応表なし', *key, lname, ldate or '', '', '', ''))

order = {'ずれ': 0, '未設定': 1, 'バーなし': 2, '対応表なし': 3, '一致': 4, '工程表なし': 5}
for o in sorted(out, key=lambda o: order[o[0]]): print(*o, sep='\t')
from collections import Counter
c = Counter(o[0] for o in out)
print('# 件数', *(f'{k}{c[k]}' for k in order if c[k]), sep='\t', file=sys.stderr)
